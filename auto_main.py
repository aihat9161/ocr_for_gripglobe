import os
import logging
from convert_to_base64 import data_to_base64
from ai_request import call_openai_api, extract_info_from_response
from openpyxl import Workbook, load_workbook
import json

# ログ設定
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
file_path = 'input_testcase/Failed2'
output_excel_path = 'output_results.xlsx'

def write_to_excel(wb, ws, filename, full_json, extracted_json, amount, date, trading_partner, invoice_registration_number):
    if ws.max_row == 1:  # If this is the first row, add headers
        headers = ["ファイル名", "JSON全て", "抽出したJSON", "amount", "date", "trading_partner", "invoice_registration_number", "全項目存在チェック"]
        ws.append(headers)
    
    all_items_present = all([amount, date, trading_partner, invoice_registration_number])
    row_data = [filename, json.dumps(full_json), json.dumps(extracted_json), amount, date, trading_partner, invoice_registration_number, "Yes" if all_items_present else "No"]
    ws.append(row_data)

def main(path):
    if os.path.exists(output_excel_path):
        wb = load_workbook(output_excel_path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active  # 新規作成したWorkbookでシートをアクティブにする

    for pathname, dirnames, filenames in os.walk(path):
        for filename in filenames:
            file = os.path.join(pathname, filename)

            try:
                supported_formats = ('.png', '.jpg', '.jpeg', '.heif', '.pdf', '.xlsx')
                
                if not filename.lower().endswith(supported_formats):
                    logging.error(f"未対応のファイル形式です: {file}")
                    continue
                
                file_size = os.path.getsize(file)
                file_format = filename.split('.')[-1]
                logging.info(f"ファイル形式：{file_format}, ファイルサイズ：{file_size} bytes")

                try:
                    with open(file, mode="rb") as f:
                        bin_file = f.read()

                    base64_data = data_to_base64(bin_file, file_format)
                    if not base64_data:
                        logging.error(f"ファイルのbase64変換に失敗しました: {file}")
                        continue
                except Exception as e:
                    logging.error(f"{file}のbase64変換中にエラーが発生しました: {e}")
                    continue

                try:
                    logging.info(f"OpenAI APIにリクエスト送信中...")
                    response = call_openai_api(base64_data[0] if isinstance(base64_data, list) else base64_data)
                    if response:
                        logging.info(f"APIリクエスト成功: {response}")
                    else:
                        logging.error(f"APIリクエスト失敗: {response}")
                except Exception as e:
                    logging.error(f"{file}のAPIリクエスト送信中にエラーが発生しました: {e}")
                    continue

                try:
                    extracted_info = extract_info_from_response(response)
                    logging.info(f"抽出された情報: {extracted_info}")
                    
                    # Write results to Excel
                    write_to_excel(
                        wb,
                        ws,
                        filename,
                        response,
                        extracted_info,
                        extracted_info.get('amount'),
                        extracted_info.get('date'),
                        extracted_info.get('trading_partner'),
                        extracted_info.get('invoice_registration_number')
                    )
                    wb.save(output_excel_path)  # 各ファイル処理ごとに保存
                except Exception as e:
                    logging.error(f"{file}のレスポンスから情報抽出中にエラーが発生しました: {e}")
                    continue

            except Exception as e:
                logging.error(f"ファイル処理中にエラーが発生しました: {e}")

    # Save the Excel workbook after all processing
    wb.save(output_excel_path)
    logging.info(f"結果を {output_excel_path} に保存しました。")

if __name__ == "__main__":
    main(file_path)
